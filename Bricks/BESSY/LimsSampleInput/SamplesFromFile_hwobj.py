#!/usr/bin/env python3
# -*- coding: utf-8 -*-

__author__  = "Thomas Hauss"
__email__ = "hauss@helmholtz-berlin.de"
__date__ = "2023-05-24 21:00 CEST"
__version__ = "0.8"

"""
Description: read a csv or Excel xlsl file with sample information.
Will put samples as list of classes into the IPSyB pipeline.

This file will be imported by Qt_SamplesFromFileWidgets

"""
#import openpyxl
import csv
try:
    import openpyxl
    xlsx = True
except:
    xlsx = False

# define the sample changer for BL-14.1 (CATS) or BL-14.2 (ISARA)
MAX_PUCKS = 30         # maximal number of pucks in a csv file
#dewar_name = "CATS"
#DEWAR_PUCKS = 9
dewar_name = "ISARA"
DEWAR_PUCKS = 29

sort_list = {}
 
class DiffractionPlan(object):
    """not used"""
    diffractionPlanId = 0 
    experimentKind = "Default"
    numberOfPositions = 0
    observedResolution = 0.0
    preferredBeamDiameter = 0.0
    radiationSensitivity = 0.0
    requiredCompleteness = 0.0
    requiredMultiplicity = 0.0
    requiredResolution = 0.0

class LimsSample(object):
    """only items used at HZB are activated by now"""
#    cellA = 0.0
#    cellAlpha = 0.0
#    cellB = 0.0
#    cellBeta = 0.0
#    cellC = 0.0
#    cellGamma = 0.0 
    proteinAcronym = ""
#    crystalSpaceGroup = "" 
#    code = ""
#    holderLength = 22.0
    containerSampleChangerLocation = -1
#    diffractionPlan = DiffractionPlan()
#    experimentType = "OSC"
#    sampleId = -1
    sampleLocation = "-1"
    sampleName = ""
    # HZB items
    containerName = ""

class ReadSamples(object):
    def __init__(self,parent):
        super(ReadSamples,self).__init__()
        self.max_pucks = MAX_PUCKS
        self.unipuck = self.sort_pucks("")

    def read_samples(self, sample_file):

        def get_delimiter(file_path, bytes = 4096):
            delimiters = (',' , ';', ':')
            sniffer = csv.Sniffer()
            data = open(file_path, "r").read(bytes)
            delimiter = sniffer.sniff(data, delimiters).delimiter
            return delimiter


        if sample_file:
            if sample_file.endswith("xlsx"):
                """will be used if python >= 3.6 and openpyxl is imported"""
                workbook = openpyxl.Workbook()
                workbook = openpyxl.load_workbook(sample_file)
                worksheet = workbook.active
                csv_worksheet = [
                    list(value) 
                    for value in worksheet.iter_rows(values_only=True)
                ]
            elif sample_file.endswith("csv"):
                delimiter = get_delimiter(sample_file)
                csv_worksheet = []
                with open(sample_file, 'r') as csvFile:
                    reader = csv.reader(csvFile, delimiter=delimiter)
                    for line in reader:
                        csv_worksheet.append(line)
            else:
                return
            return csv_worksheet         

    def resort_puck_position(self, samples, sort_list):
        samples_list = []
        for sample in samples:
            for key, value in sort_list.items():
                for name, position in value.items():
                    if sample.containerName == name:
                        sample.containerSampleChangerLocation = str(position)
        for sample in samples:
            if int(sample.containerSampleChangerLocation) > 0:
                samples_list.append(sample)
        return samples_list
                        
    def sort_pucks(self,csv_worksheet = ""):
        # Definition of the column entries in the csv sample file, according to MAX-IV
        # line[0] = NameDewar     # not used
        # line[1] = PuckName
        # line[2] = PuckType
        # line[3] = SamplePosition
        # line[4] = ProteinAcronym
        # line[5] = SampleName
        # line[6] = Pin Barcode     # not used
        
        pucknames = [''] * MAX_PUCKS
        number_of_pucks = 0
        unipuck = {}
        dewar = {}
        current = ""
        i = 1
        j = 1
        if csv_worksheet != "":        
            for line in csv_worksheet:
                if line[0].startswith('#') or len(line) < 2:
                    continue
                else:
                    name = line[1]
                    if  not any(name in word for word in pucknames):
                        pucknames[number_of_pucks] = name
                        number_of_pucks += 1
        return pucknames, number_of_pucks

    def csv_worksheet_to_samples(self,csv_worksheet):
        """Sort csv samples into lims_samples class. Attributes are defined in LimsSample"""
        samples = []
        for line in csv_worksheet:
            if line[0].startswith('#') or len(line) < 2:
                continue
            else:
                lims_sample = LimsSample()
                if hasattr(lims_sample, 'cellA'):
                    lims_sample.cellA = 0.0
                if hasattr(lims_sample, 'cellAalpha'):
                    lims_sample.cellAlpha = 0.0
                if hasattr(lims_sample, 'cellB'):
                    lims_sample.cellB = 0.0
                if hasattr(lims_sample, 'cellBeta'):
                    lims_sample.cellBeta = 0.0
                if hasattr(lims_sample, 'cellC'):
                    lims_sample.cellC = 0.0
                if hasattr(lims_sample, 'cellGamma'):
                    lims_sample.cellGamma = 0.0 
                if hasattr(lims_sample, 'containerSampleChangerLocation'):
                    lims_sample.containerSampleChangerLocation = "0"
                if hasattr(lims_sample, 'crystalSpaceGroup'):
                    lims_sample.crystalSpaceGroup = "" 
                if hasattr(lims_sample, 'diffractionPlan'):
                    lims_sample.diffractionPlan
                if hasattr(lims_sample, 'diffractionPlan.diffractionPlanId'):
                    lims_sample.diffractionPlan.diffractionPlanId = 0 
                if hasattr(lims_sample, 'diffractionPlan.experimentKind'):
                    lims_sample.diffractionPlan.experimentKind = "Default"
                if hasattr(lims_sample, 'diffractionPlan.numberOfPositions'):
                    lims_sample.diffractionPlan.numberOfPositions = 0
                if hasattr(lims_sample, 'diffractionPlan.observedResolution'):
                    lims_sample.diffractionPlan.observedResolution = 0.0
                if hasattr(lims_sample, 'diffractionPlan.preferredBeamDiameter'):
                    lims_sample.diffractionPlan.preferredBeamDiameter = 0.0
                if hasattr(lims_sample, 'diffractionPlan.radiationSensitivity'):
                    lims_sample.diffractionPlan.radiationSensitivity = 0.0
                if hasattr(lims_sample, 'diffractionPlan.requiredCompleteness'):
                    lims_sample.diffractionPlan.requiredCompleteness = 0.0
                if hasattr(lims_sample, 'diffractionPlan.requiredMultiplicity'):
                    lims_sample.diffractionPlan.requiredMultiplicity = 0.0
                if hasattr(lims_sample, 'diffractionPlan.requiredResolution'):
                    lims_sample.diffractionPlan.requiredResolution = 0.0
                if hasattr(lims_sample, 'experimentType'):
                    lims_sample.experimentType = "OSC"
                if hasattr(lims_sample, 'proteinAcronym'):
                    lims_sample.proteinAcronym = line[4]
                if hasattr(lims_sample, 'sampleId'):
                    lims_sample.sampleId = 0
                if hasattr(lims_sample, 'sampleLocation'):
                    lims_sample.sampleLocation = line[3]
                if hasattr(lims_sample, 'sampleName'):
                    lims_sample.sampleName = line[5]
                # HZB items
                if hasattr(lims_sample, 'containerName'):
                    lims_sample.containerName = line[1]
                samples.append(lims_sample)
        return samples